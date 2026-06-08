"""
@file merge.py
@brief DOCX merge utility combining base template with generated body document.

Combines:
- base DOCX (template + metadata)
- rendered body DOCX (Pandoc output)

Preserves formatting via docxcompose.
"""

import sys
import os
from docx import Document
from docxcompose.composer import Composer
from docxtpl import DocxTemplate

# ---------------- Argument Parsing ----------------
# Parse CLI arguments: env file, base document, body document, output path
env_file = sys.argv[1]
base_doc = sys.argv[2]
body_doc = sys.argv[3]
output_doc = sys.argv[4]

# ---------------- Environment Loading ----------------
# Load variables from .env file
context = {}
if os.path.exists(env_file):
    with open(env_file, "r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            context[key.strip()] = value.strip().strip('"').strip("'")

# ---------------- Base Document Rendering ----------------
# Render base DOCX placeholders like {{NAME}}
rendered_base = base_doc + ".rendered.docx"
tpl = DocxTemplate(base_doc)
tpl.render(context)
tpl.save(rendered_base)

# ---------------- Document Composition ----------------
# Load base DOCX as master document
master = Document(rendered_base)
# Initialize DOCX composer for merging documents
composer = Composer(master)

# Append generated body document to base document
composer.append(Document(body_doc))

# ---------------- Output ----------------
# Save final merged DOCX file
composer.save(output_doc)